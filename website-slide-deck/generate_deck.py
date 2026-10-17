"""
Generates editable 16:9 PPTX mirroring Ly Thanh Loi portfolio (current codebase).
Run: pip install python-pptx && python generate_deck.py
"""
from __future__ import annotations

import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

OUT_DIR = Path(__file__).resolve().parent
PPTX_PATH = OUT_DIR / "Portfolio_Current_Version.pptx"

# Verified from src/app/globals.css
C = {
    "void": RGBColor(0x07, 0x08, 0x0F),
    "void_elevated": RGBColor(0x0E, 0x10, 0x19),
    "panel": RGBColor(0x12, 0x15, 0x1F),
    "border": RGBColor(0x1E, 0x24, 0x33),
    "cyan": RGBColor(0x22, 0xD3, 0xEE),
    "violet": RGBColor(0xA7, 0x8B, 0xFA),
    "text": RGBColor(0xE8, 0xEA, 0xF0),
    "muted": RGBColor(0x94, 0x99, 0xAB),
    "amber": RGBColor(0xFB, 0xBF, 0x24),
    "emerald": RGBColor(0x34, 0xD3, 0x99),
    "overlay": RGBColor(0x00, 0x00, 0x00),
}

# Font substitutes (web: Cinzel, DM Sans, JetBrains Mono)
FONT_DISPLAY = "Georgia"
FONT_BODY = "Calibri"
FONT_MONO = "Consolas"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

NAV_LINKS = ["Entrance", "Profile", "Abilities", "Archive", "Timeline", "Portal"]

PROFILE = {
    "name": "Ly Thanh Loi",
    "mono_tag": "digital_grimoire.init()",
    "headline": "Information Technology student bridging infrastructure, analytical thinking, and Generative AI to build reliable software systems.",
    "career": "Leveraging IT infrastructure, analytical thinking, and Generative AI to optimize operational processes and data — creating tangible value while growing skills in coordinating comprehensive software systems.",
    "intro": "Fourth-year student at HUTECH (Vietnam–Japan Institute). I combine backend development, cross-platform apps, and AI integration with a focus on accurate data flow, standardized processes, and operational performance.",
    "direction": "Building a solid foundation in system operations and administration, combining development logic with enterprise-oriented solutions and cloud infrastructure.",
    "email": "ltl2004123@gmail.com",
    "github": "@LLLVPN4423",
}

ATTRIBUTES = [
    ("Identity", "Ly Thanh Loi"),
    ("Institute", "Information Technology — Vietnam–Japan Institute"),
    ("Academy", "Ho Chi Minh City University of Technology (HUTECH)"),
    ("Cycle", "2022 – 2026"),
    ("Status", "Fourth-year student"),
    ("Focus", "Backend Developer · Team Lead · UI/UX Optimization"),
]

SKILL_GROUPS = [
    ("◈", "Programming Languages", "C# · Java · JavaScript"),
    ("⬡", "Frameworks & Platforms", ".NET MAUI · ASP.NET Core · Spring Boot (J2EE) · WinForms"),
    ("◇", "Databases", "SQL Server · MySQL · SQLite · Firebase"),
    ("✦", "Architecture & Patterns", "MVVM · Repository Pattern · Dependency Injection · RESTful API"),
    ("◎", "Testing & Quality", "NUnit · Gauge Framework"),
    ("⎔", "Tools, AI & Platforms", "Git / GitHub · Adobe XD · GitHub Copilot · Gemini, ChatGPT, Claude · Google Workspace"),
]

SKILL_DETAIL = ("Programming Languages", ["C# (.NET 9 / 10)", "Java", "JavaScript (React)"])

PROJECTS = [
    ("violet", "✎", "CrayonNote", "Cross-platform note-taking app (.NET MAUI) with AI-powered speech-to-text.", ".NET MAUI", "C#", "SQLite (ORM)", "OpenAI Whisper API"),
    ("cyan", "文", "MiniHongo", "Mobile app for Japanese vocabulary and Kanji learning across platforms.", ".NET MAUI", "C#", "SQLite", "RESTful API"),
    ("amber", "☽", "Bé Tabby Tarot (Pet Tarot)", "Bilingual (VI/EN) pet-themed tarot web app — 22 Major Arcana with shuffle and journal.", "TypeScript", "React", "Next.js"),
    ("emerald", "☁", "AWS Internship Report Site", "Bilingual Hugo site documenting AWS Vietnam FCJ Cloud Engineer internship.", "Hugo", "GitHub Actions", "GitHub Pages"),
    ("cyan", "☑", "Task Management System", "Online task management and team collaboration with permissions and progress tracking.", "Java Spring Boot", "MySQL", "Spring Security", "RESTful API"),
    ("violet", "♫", "Music Player (Desktop)", "Windows multimedia player with modern UI, playlists, and multiple formats.", "C#", "WinForms", ".NET", "NUnit"),
]

CRAYON_MODAL = {
    "purpose": "Optimize digitization of audio recordings and multimedia data management on mobile devices.",
    "role": "End-to-end development",
    "features": [
        "Cross-platform .NET MAUI client",
        "Local data with SQLite ORM",
        "Speech-to-text via OpenAI Whisper API",
        "Azure OpenAI integration",
        "MVVM architecture",
    ],
    "tech": [".NET MAUI", "C#", "SQLite (ORM)", "OpenAI Whisper API", "Azure OpenAI", "MVVM"],
}

TIMELINE = [
    ("2022 – 2026 · Academy", "Information Technology", "HUTECH — Vietnam–Japan Institute", "Fourth-year student; focus on software systems, backend, and cross-platform development."),
    ("04/2024 – 06/2024 · Field Ops", "Technician", "Viet Nghiep Electronics Technology Co., Ltd.", "Manufacturing circuit boards, programming, and completing electronic products."),
    ("2024 – Present · Field Ops", "Tour Guide", "APEX Vietnam Travel Corporation", "Lead student groups; communicate in Japanese and English per project schedule."),
    ("— · Field Ops", "Personal Trainer", "TODAY FITNESS", "Plan training schedules, support trainees, and digitize training progress records."),
    ("04/2026 – 07/2026 · Signal", "FCJ Cloud Engineer Intern", "AWS Vietnam — First Cloud AI Journey", "Hands-on cloud training; network infrastructure, servers, large-scale databases; Jira task management and data-flow optimization."),
    ("03/2026 · Signal", 'Webinar "Beyond Code"', "FPT Software", "Enterprise solutions and low-code trends; leadership beyond pure technology."),
    ("10/2025 · Signal", "TechTogether", "Bosch Vietnam", "Industry 4.0 experience and career guidance from experts."),
]

SLIDE_MANIFEST: list[dict] = []


def rgb_fill(shape, color: RGBColor, transparency: float = 0.0):
    fill = shape.fill
    fill.solid()
    fill.fore_color.rgb = color
    if transparency:
        fill.transparency = transparency


def set_text(
    shape,
    text: str,
    size: float,
    color: RGBColor,
    bold: bool = False,
    font: str = FONT_BODY,
    align=PP_ALIGN.LEFT,
):
    tf = shape.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = font


def add_labeled_note(slide, label: str, slide_id: str):
    box = slide.shapes.add_textbox(Inches(0.25), Inches(7.05), Inches(12.8), Inches(0.35))
    set_text(box, f"[{slide_id}] {label} — editable mock; fonts: Georgia≈Cinzel, Calibri≈DM Sans, Consolas≈JetBrains Mono", 8, C["muted"], font=FONT_MONO)
    box.name = f"meta_{slide_id}"


def add_viewport_bg(slide, with_grid: bool = True):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg.name = "bg_void"
    rgb_fill(bg, C["void"])
    bg.line.fill.background()

    frame = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.35), Inches(0.2), Inches(12.63), Inches(6.75))
    frame.name = "viewport_frame"
    rgb_fill(frame, C["void"])
    frame.line.color.rgb = C["border"]

    if with_grid:
        for i in range(0, 14):
            x = Inches(0.35 + i * 0.45)
            ln = slide.shapes.add_connector(1, x, Inches(0.55), x, Inches(6.7))
            ln.line.color.rgb = C["border"]
            ln.line.width = Pt(0.35)
            ln.name = f"grid_v_{i}"
        for j in range(0, 8):
            y = Inches(0.55 + j * 0.45)
            ln = slide.shapes.add_connector(1, Inches(0.35), y, Inches(12.98), y)
            ln.line.color.rgb = C["border"]
            ln.line.width = Pt(0.35)
            ln.name = f"grid_h_{j}"


def add_scroll_progress(slide, width_in: float):
    track = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.35), Inches(0.2), Inches(12.63), Inches(0.04))
    track.name = "scroll_track"
    rgb_fill(track, C["void_elevated"])
    track.line.fill.background()
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.35), Inches(0.2), Inches(width_in), Inches(0.04))
    bar.name = "scroll_progress_bar"
    rgb_fill(bar, C["cyan"])
    bar.line.fill.background()


def add_nav(slide, active: str, mobile: bool = False, menu_open: bool = False):
    header = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.35), Inches(0.24), Inches(12.63), Inches(0.55))
    header.name = "nav_header"
    rgb_fill(header, C["void"], transparency=0.2)
    header.line.color.rgb = C["border"]

    logo = slide.shapes.add_textbox(Inches(0.55), Inches(0.34), Inches(2.2), Inches(0.35))
    logo.name = "nav_logo"
    set_text(logo, "LTL ARCHIVE", 11, C["cyan"], font=FONT_DISPLAY)

    if mobile and not menu_open:
        menu_btn = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(11.85), Inches(0.34), Inches(0.85), Inches(0.32))
        menu_btn.name = "nav_menu_button"
        rgb_fill(menu_btn, C["void"])
        menu_btn.line.color.rgb = C["border"]
        set_text(menu_btn, "Menu", 9, C["muted"], align=PP_ALIGN.CENTER)
    elif not mobile:
        x = Inches(3.2)
        for link in NAV_LINKS:
            pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(0.34), Inches(1.05), Inches(0.32))
            pill.name = f"nav_link_{link.lower()}"
            if link == active:
                rgb_fill(pill, C["panel"])
                set_text(pill, link.upper(), 8, C["cyan"], align=PP_ALIGN.CENTER, font=FONT_BODY)
            else:
                rgb_fill(pill, C["void"])
                pill.line.fill.background()
                set_text(pill, link.upper(), 8, C["muted"], align=PP_ALIGN.CENTER, font=FONT_BODY)
            x += Inches(1.12)

    if mobile and menu_open:
        drawer = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.35), Inches(0.79), Inches(12.63), Inches(1.55))
        drawer.name = "nav_mobile_drawer"
        rgb_fill(drawer, C["void_elevated"])
        drawer.line.color.rgb = C["border"]
        y = Inches(0.9)
        for link in NAV_LINKS:
            tb = slide.shapes.add_textbox(Inches(0.65), y, Inches(3), Inches(0.22))
            tb.name = f"mobile_nav_{link.lower()}"
            set_text(tb, link, 11, C["text"])
            y += Inches(0.22)


def section_heading(slide, top: float, eyebrow: str, title: str, eyebrow_color: RGBColor):
    e = slide.shapes.add_textbox(Inches(0.65), Inches(top), Inches(5), Inches(0.25))
    e.name = "section_eyebrow"
    set_text(e, eyebrow.upper(), 9, eyebrow_color, font=FONT_MONO)
    t = slide.shapes.add_textbox(Inches(0.65), Inches(top + 0.28), Inches(8), Inches(0.45))
    t.name = "section_title"
    set_text(t, title, 28, C["text"], font=FONT_DISPLAY)


def slide_hero_desktop(prs, slide_id: str, note: str):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_viewport_bg(slide)
    add_scroll_progress(slide, 1.2)
    add_nav(slide, active="Entrance")
    add_labeled_note(slide, note, slide_id)

    glow = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(7.5), Inches(2.2), Inches(2.2), Inches(2.2))
    glow.name = "hero_violet_glow"
    rgb_fill(glow, C["violet"], transparency=0.85)
    glow.line.fill.background()

    tag = slide.shapes.add_textbox(Inches(0.65), Inches(1.15), Inches(4), Inches(0.2))
    set_text(tag, PROFILE["mono_tag"].upper(), 9, C["cyan"], font=FONT_MONO)

    title = slide.shapes.add_textbox(Inches(0.65), Inches(1.45), Inches(6.5), Inches(0.7))
    title.name = "hero_name"
    set_text(title, PROFILE["name"], 36, C["text"], font=FONT_DISPLAY)

    h = slide.shapes.add_textbox(Inches(0.65), Inches(2.25), Inches(6.2), Inches(0.9))
    h.name = "hero_headline"
    set_text(h, PROFILE["headline"], 14, C["muted"])

    c = slide.shapes.add_textbox(Inches(0.65), Inches(3.15), Inches(6.2), Inches(0.75))
    c.name = "hero_career_goal"
    set_text(c, PROFILE["career"], 11, C["muted"])

    for i, (label, primary) in enumerate([("Explore My World", True), ("View Projects", False), ("GitHub", False)]):
        btn = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.65 + i * 1.85), Inches(4.05), Inches(1.65), Inches(0.38))
        btn.name = f"hero_btn_{label.replace(' ', '_').lower()}"
        rgb_fill(btn, C["panel"] if primary else C["void"])
        btn.line.color.rgb = C["border"] if not primary else C["cyan"]
        set_text(btn, label, 10, C["text"] if primary else C["muted"], align=PP_ALIGN.CENTER)

    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.2), Inches(1.5), Inches(3.8), Inches(3.8))
    card.name = "hero_portal_card"
    rgb_fill(card, C["void_elevated"])
    card.line.color.rgb = C["border"]

    inner = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.55), Inches(1.85), Inches(3.1), Inches(3.1))
    inner.name = "hero_portal_dashed"
    rgb_fill(inner, C["void_elevated"])
    inner.line.color.rgb = C["violet"]
    inner.line.dash_style = 2

    star = slide.shapes.add_textbox(Inches(9.35), Inches(2.5), Inches(1.5), Inches(0.6))
    set_text(star, "✦", 40, C["violet"], align=PP_ALIGN.CENTER, font=FONT_DISPLAY)
    core = slide.shapes.add_textbox(Inches(8.7), Inches(3.2), Inches(3.5), Inches(0.5))
    set_text(core, "CORE MODULE\nFantasy × Tech Portal", 10, C["cyan"], align=PP_ALIGN.CENTER, font=FONT_MONO)
    meta = slide.shapes.add_textbox(Inches(8.55), Inches(4.0), Inches(3.1), Inches(0.9))
    set_text(meta, "status: online\nfocus: systems + AI\narchive: LLLVPN4423", 9, C["muted"], align=PP_ALIGN.CENTER, font=FONT_MONO)

    SLIDE_MANIFEST.append({"id": slide_id, "type": "hero", "viewport": "desktop", "verified": "code+production URL"})


def slide_about(prs, slide_id: str):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_viewport_bg(slide, with_grid=False)
    add_scroll_progress(slide, 4.5)
    add_nav(slide, active="Profile")
    add_labeled_note(slide, "About Me — Character Profile", slide_id)
    section_heading(slide, 1.05, "Character Profile", "About Me", C["violet"])

    left = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.65), Inches(1.85), Inches(5.9), Inches(3.2))
    left.name = "about_summary_card"
    rgb_fill(left, C["panel"], transparency=0.5)
    left.line.color.rgb = C["border"]
    set_text(left, "SUMMARY\n\n" + PROFILE["intro"] + "\n\n" + PROFILE["direction"], 11, C["muted"])

    right = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.75), Inches(1.85), Inches(5.9), Inches(3.2))
    right.name = "about_attributes_card"
    rgb_fill(right, C["void_elevated"])
    right.line.color.rgb = C["border"]
    lines = "ATTRIBUTES\n\n" + "\n".join(f"{k}: {v}" for k, v in ATTRIBUTES)
    set_text(right, lines, 10, C["text"])

    SLIDE_MANIFEST.append({"id": slide_id, "type": "about", "verified": "code"})


def slide_skills(prs, slide_id: str, expanded: bool):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_viewport_bg(slide, with_grid=False)
    add_scroll_progress(slide, 3.8)
    add_nav(slide, active="Abilities")
    add_labeled_note(slide, "Skills — module detail open" if expanded else "Skills — default", slide_id)
    section_heading(slide, 1.05, "Technical Abilities", "Skill Modules", C["cyan"])
    sub = slide.shapes.add_textbox(Inches(0.65), Inches(1.72), Inches(10), Inches(0.3))
    set_text(sub, "Verified skills from my profile — select a module to inspect nodes.", 11, C["muted"])

    positions = [(0.65, 2.15), (4.55, 2.15), (8.45, 2.15), (0.65, 3.55), (4.55, 3.55), (8.45, 3.55)]
    for idx, ((x, y), (rune, title, preview)) in enumerate(zip(positions, SKILL_GROUPS)):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(3.65), Inches(1.15))
        card.name = f"skill_card_{idx}"
        rgb_fill(card, C["panel"], transparency=0.6)
        border = C["cyan"] if expanded and title == SKILL_DETAIL[0] else C["border"]
        card.line.color.rgb = border
        set_text(card, f"{rune}\n{title}\n{preview}", 9, C["text"])

    if expanded:
        panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.65), Inches(4.95), Inches(11.95), Inches(1.35))
        panel.name = "skill_detail_panel"
        rgb_fill(panel, C["void_elevated"])
        panel.line.color.rgb = C["cyan"]
        chips = "  ".join(f"[{t}]" for t in SKILL_DETAIL[1])
        set_text(panel, f"MODULE DETAIL — {SKILL_DETAIL[0]}\nDismiss\n\n{chips}", 10, C["text"])

    SLIDE_MANIFEST.append({"id": slide_id, "type": "skills", "state": "expanded" if expanded else "default", "verified": "code"})


def accent_rgb(name: str) -> RGBColor:
    return {"cyan": C["cyan"], "violet": C["violet"], "amber": C["amber"], "emerald": C["emerald"]}[name]


def add_project_card(slide, x, y, w, h, accent, glyph, name, desc, tags, hover: bool = False):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    card.name = f"project_card_{name.replace(' ', '_')[:20]}"
    rgb_fill(card, C["panel"], transparency=0.7)
    card.line.color.rgb = C["cyan"] if hover else C["border"]
    grad = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h * 0.55))
    grad.name = f"project_card_gradient_{name[:12]}"
    rgb_fill(grad, accent_rgb(accent), transparency=0.85)
    grad.line.fill.background()
    body = slide.shapes.add_textbox(Inches(x + 0.15), Inches(y + 0.15), Inches(w - 0.3), Inches(h - 0.55))
    set_text(body, f"{glyph}                    OPEN RECORD\n{name}\n{desc}\n\n{' · '.join(tags)}", 9, C["text"])
    footer = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y + h - 0.35), Inches(w), Inches(0.35))
    footer.name = f"project_card_footer_{name[:12]}"
    rgb_fill(footer, C["void"], transparency=0.3)
    footer.line.color.rgb = C["border"]
    links = slide.shapes.add_textbox(Inches(x + 0.15), Inches(y + h - 0.3), Inches(w), Inches(0.25))
    set_text(links, "GitHub" + ("   Live Demo" if "AWS" in name else ""), 9, C["cyan"])


def slide_projects(prs, slide_id: str, start: int, hover_index: int | None = None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_viewport_bg(slide, with_grid=False)
    add_scroll_progress(slide, 5.2)
    add_nav(slide, active="Archive")
    add_labeled_note(slide, f"Projects grid cards {start + 1}–{min(start + 4, 6)}", slide_id)
    section_heading(slide, 1.05, "Project Archive", "Verified Artifacts", C["violet"])
    sub = slide.shapes.add_textbox(Inches(0.65), Inches(1.72), Inches(10), Inches(0.35))
    set_text(sub, "Projects sourced from my CV and public GitHub repositories — each links to the original repo.", 11, C["muted"])

    batch = PROJECTS[start : start + 4]
    coords = [(0.65, 2.2), (6.75, 2.2), (0.65, 4.05), (6.75, 4.05)]
    for i, ((x, y), proj) in enumerate(zip(coords, batch)):
        accent, glyph, name, desc, *tags = proj
        global_idx = start + i
        add_project_card(slide, x, y, 5.85, 1.65, accent, glyph, name, desc, tags[:4], hover=(hover_index == global_idx))

    SLIDE_MANIFEST.append({"id": slide_id, "type": "projects", "cards": [p[2] for p in batch], "verified": "code"})


def slide_project_modal(prs, slide_id: str, dimmed: bool):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_viewport_bg(slide, with_grid=False)
    add_scroll_progress(slide, 5.2)
    add_nav(slide, active="Archive")
    add_labeled_note(slide, "Project detail modal — CrayonNote (open)" if not dimmed else "Before open — projects visible", slide_id)

    if dimmed:
        add_project_card(slide, 0.65, 2.2, 5.85, 1.65, "violet", "✎", "CrayonNote", PROJECTS[0][3], list(PROJECTS[0][4:8]))
        hint = slide.shapes.add_textbox(Inches(7), Inches(3), Inches(5), Inches(0.4))
        set_text(hint, "User clicks card → modal opens (see next slide)", 11, C["muted"])
    else:
        overlay = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.35), Inches(0.24), Inches(12.63), Inches(6.71))
        overlay.name = "modal_overlay"
        rgb_fill(overlay, C["overlay"], transparency=0.3)
        overlay.line.fill.background()

        modal = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2.4), Inches(0.85), Inches(8.5), Inches(5.85))
        modal.name = "project_modal_panel"
        rgb_fill(modal, C["void_elevated"])
        modal.line.color.rgb = C["violet"]

        header = slide.shapes.add_textbox(Inches(2.65), Inches(1.0), Inches(7.5), Inches(0.55))
        set_text(header, "ARTIFACT RECORD\n✎  CrayonNote\nClose", 11, C["text"], font=FONT_MONO)

        content = slide.shapes.add_textbox(Inches(2.65), Inches(1.65), Inches(7.9), Inches(4.7))
        feat = "\n".join(f"• {f}" for f in CRAYON_MODAL["features"])
        tech = "  ".join(f"[{t}]" for t in CRAYON_MODAL["tech"])
        txt = (
            f"OVERVIEW\n{PROJECTS[0][3]}\n\n"
            f"PROBLEM & PURPOSE\n{CRAYON_MODAL['purpose']}\n\n"
            f"CONTRIBUTION\n{CRAYON_MODAL['role']}\n\n"
            f"FEATURES\n{feat}\n\n"
            f"TECHNOLOGY STACK\n{tech}\n\n"
            f"View Source (GitHub)"
        )
        set_text(content, txt, 10, C["muted"])

    SLIDE_MANIFEST.append({"id": slide_id, "type": "project_modal", "project": "CrayonNote", "state": "before" if dimmed else "open", "verified": "code"})


def slide_timeline(prs, slide_id: str, start: int, count: int):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_viewport_bg(slide, with_grid=False)
    add_scroll_progress(slide, 6.8)
    add_nav(slide, active="Timeline")
    add_labeled_note(slide, f"Timeline scroll continuation entries {start + 1}+", slide_id)
    section_heading(slide, 1.05, "Experience & Timeline", "Chronicle", C["cyan"])

    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.85), Inches(1.85), Inches(0.03), Inches(4.8))
    line.name = "timeline_axis"
    rgb_fill(line, C["border"])
    line.line.fill.background()

    y = 1.95
    for entry in TIMELINE[start : start + count]:
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.78), Inches(y), Inches(0.12), Inches(0.12))
        dot.name = "timeline_dot"
        rgb_fill(dot, C["violet"])
        dot.line.fill.background()
        tb = slide.shapes.add_textbox(Inches(1.15), Inches(y - 0.05), Inches(11), Inches(0.85))
        period, title, org, desc = entry
        set_text(tb, f"{period}\n{title}\n{org}\n{desc}", 10, C["muted"])
        y += 0.95

    SLIDE_MANIFEST.append({"id": slide_id, "type": "timeline", "entries": count, "verified": "code"})


def slide_contact_footer(prs, slide_id: str):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_viewport_bg(slide, with_grid=False)
    add_scroll_progress(slide, 11.5)
    add_nav(slide, active="Portal")
    add_labeled_note(slide, "Contact + Footer", slide_id)
    section_heading(slide, 1.05, "The Portal", "Connect", C["violet"])
    sub = slide.shapes.add_textbox(Inches(0.65), Inches(1.72), Inches(8), Inches(0.3))
    set_text(sub, "Open a channel for collaboration, internships, or project discussions.", 11, C["muted"])

    for i, (label, value, color) in enumerate([
        ("Email", PROFILE["email"], C["cyan"]),
        ("GitHub", PROFILE["github"], C["text"]),
    ]):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.65 + i * 3.2), Inches(2.2), Inches(2.95), Inches(1.1))
        card.name = f"contact_card_{label.lower()}"
        rgb_fill(card, C["panel"], transparency=0.5)
        card.line.color.rgb = C["border"]
        set_text(card, f"{label.upper()}\n{value}", 11, color)

    foot = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.35), Inches(6.35), Inches(12.63), Inches(0.55))
    foot.name = "footer_bar"
    rgb_fill(foot, C["void"])
    foot.line.color.rgb = C["border"]
    set_text(foot, f"© 2026 Ly Thanh Loi — Digital Archive          Projects    GitHub    Back to top", 10, C["muted"], align=PP_ALIGN.LEFT)

    SLIDE_MANIFEST.append({"id": slide_id, "type": "contact_footer", "verified": "code"})


def slide_mobile_hero(prs, slide_id: str):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_viewport_bg(slide)
    add_scroll_progress(slide, 0.8)
    add_nav(slide, active="Entrance", mobile=True, menu_open=False)
    add_labeled_note(slide, "Responsive — mobile hero (< md breakpoint)", slide_id)

    tag = slide.shapes.add_textbox(Inches(0.65), Inches(1.05), Inches(5), Inches(0.2))
    set_text(tag, PROFILE["mono_tag"].upper(), 9, C["cyan"], font=FONT_MONO)
    title = slide.shapes.add_textbox(Inches(0.65), Inches(1.35), Inches(11.5), Inches(0.55))
    set_text(title, PROFILE["name"], 30, C["text"], font=FONT_DISPLAY)
    h = slide.shapes.add_textbox(Inches(0.65), Inches(2.0), Inches(11.5), Inches(0.8))
    set_text(h, PROFILE["headline"], 13, C["muted"])

    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2.5), Inches(3.0), Inches(8.3), Inches(2.8))
    card.name = "mobile_hero_portal_card"
    rgb_fill(card, C["void_elevated"])
    card.line.color.rgb = C["border"]
    set_text(card, "✦\nCORE MODULE\nFantasy × Tech Portal", 14, C["violet"], align=PP_ALIGN.CENTER)

    SLIDE_MANIFEST.append({"id": slide_id, "type": "responsive_mobile_hero", "verified": "code (md: breakpoints)"})


def slide_mobile_menu(prs, slide_id: str):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_viewport_bg(slide, with_grid=False)
    add_nav(slide, active="Entrance", mobile=True, menu_open=True)
    add_labeled_note(slide, "Responsive — mobile nav open", slide_id)
    SLIDE_MANIFEST.append({"id": slide_id, "type": "responsive_mobile_nav", "verified": "code"})


def slide_mobile_modal(prs, slide_id: str):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_viewport_bg(slide, with_grid=False)
    add_nav(slide, active="Archive", mobile=True)
    add_labeled_note(slide, "Responsive — modal items-end (bottom sheet on mobile)", slide_id)
    overlay = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.35), Inches(0.24), Inches(12.63), Inches(6.71))
    rgb_fill(overlay, C["overlay"], transparency=0.3)
    overlay.line.fill.background()
    modal = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.35), Inches(3.2), Inches(12.63), Inches(3.75))
    modal.name = "mobile_modal_sheet"
    rgb_fill(modal, C["void_elevated"])
    modal.line.color.rgb = C["violet"]
    set_text(modal, "ARTIFACT RECORD — CrayonNote\n(rounded-t-2xl, full width)\n…", 11, C["text"])
    SLIDE_MANIFEST.append({"id": slide_id, "type": "responsive_mobile_modal", "verified": "code (sm:items-center)"})


def slide_storyboard_reveal(prs, slide_id: str):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    rgb_fill(bg, C["void"])
    bg.line.fill.background()
    add_labeled_note(slide, "Animation storyboard — Reveal on scroll (3 frames)", slide_id)
    title = slide.shapes.add_textbox(Inches(0.5), Inches(0.35), Inches(12), Inches(0.4))
    set_text(title, "Storyboard: Section Reveal (IntersectionObserver)", 18, C["text"], font=FONT_DISPLAY)

    frames = [
        ("Frame A — Initial", "opacity: 0\ntranslateY: 20px\n(reveal-hidden)", C["border"]),
        ("Frame B — In view", "Observer threshold 0.12\nrootMargin bottom -8%", C["cyan"]),
        ("Frame C — Visible", "opacity: 1\ntranslateY: 0\n650ms cubic-bezier(0.22,1,0.36,1)", C["violet"]),
    ]
    for i, (label, body, col) in enumerate(frames):
        x = 0.65 + i * 4.1
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.2), Inches(3.7), Inches(4.5))
        box.name = f"storyboard_reveal_{i}"
        rgb_fill(box, C["panel"])
        box.line.color.rgb = col
        set_text(box, f"{label}\n\n{body}", 11, C["text"])
    SLIDE_MANIFEST.append({"id": slide_id, "type": "storyboard_reveal", "verified": "code globals.css + Reveal.tsx"})


def slide_storyboard_modal(prs, slide_id: str):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    rgb_fill(bg, C["void"])
    bg.line.fill.background()
    add_labeled_note(slide, "Animation storyboard — Project modal open/close", slide_id)
    title = slide.shapes.add_textbox(Inches(0.5), Inches(0.35), Inches(12), Inches(0.4))
    set_text(title, "Storyboard: Project Detail Panel", 18, C["text"], font=FONT_DISPLAY)
    steps = [
        "1. Click project card → setSelected(project)",
        "2. Overlay bg-black/70 covers viewport",
        "3. Panel slides up (mobile) / centered (sm+)",
        "4. body overflow hidden; Escape → onClose",
        "5. Click overlay → onClose",
    ]
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.65), Inches(1.2), Inches(11.9), Inches(4.8))
    set_text(box, "\n".join(steps), 14, C["muted"])
    SLIDE_MANIFEST.append({"id": slide_id, "type": "storyboard_modal", "verified": "code ProjectDetailPanel.tsx"})


def slide_storyboard_scroll_progress(prs, slide_id: str):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    rgb_fill(bg, C["void"])
    bg.line.fill.background()
    add_labeled_note(slide, "Scroll progress bar width vs scrollY", slide_id)
    title = slide.shapes.add_textbox(Inches(0.5), Inches(0.35), Inches(12), Inches(0.4))
    set_text(title, "Storyboard: Scroll Progress Indicator", 18, C["text"], font=FONT_DISPLAY)
    for i, w in enumerate([1.0, 4.0, 8.0]):
        y = 1.3 + i * 1.6
        track = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.65), Inches(y), Inches(11.9), Inches(0.05))
        rgb_fill(track, C["void_elevated"])
        track.line.fill.background()
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.65), Inches(y), Inches(w), Inches(0.05))
        rgb_fill(bar, C["cyan"])
        bar.line.fill.background()
        cap = slide.shapes.add_textbox(Inches(0.65), Inches(y + 0.12), Inches(10), Inches(0.35))
        set_text(cap, f"scroll progress ≈ {(w/11.9)*100:.0f}% — CSS transition width 150ms", 12, C["muted"])
    SLIDE_MANIFEST.append({"id": slide_id, "type": "storyboard_scroll_progress", "verified": "code ScrollProgress.tsx"})


def main():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    slide_hero_desktop(prs, "S01", "Initial load — Hero viewport (desktop, nav Entrance active)")
    slide_about(prs, "S02")
    slide_skills(prs, "S03", expanded=False)
    slide_skills(prs, "S04", expanded=True)
    slide_projects(prs, "S05", start=0, hover_index=None)
    slide_projects(prs, "S06", start=4, hover_index=None)
    slide_projects(prs, "S07", start=0, hover_index=0)
    slide_project_modal(prs, "S08", dimmed=True)
    slide_project_modal(prs, "S09", dimmed=False)
    slide_timeline(prs, "S10", start=0, count=4)
    slide_timeline(prs, "S11", start=4, count=3)
    slide_contact_footer(prs, "S12")
    slide_mobile_hero(prs, "S13")
    slide_mobile_menu(prs, "S14")
    slide_mobile_modal(prs, "S15")
    slide_storyboard_reveal(prs, "S16")
    slide_storyboard_modal(prs, "S17")
    slide_storyboard_scroll_progress(prs, "S18")

    prs.save(PPTX_PATH)
    (OUT_DIR / "slide_manifest.json").write_text(json.dumps(SLIDE_MANIFEST, indent=2), encoding="utf-8")
    print(f"Wrote {PPTX_PATH} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
